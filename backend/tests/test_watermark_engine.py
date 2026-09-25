"""Watermark engine on real sample files (spec 23, 24, 28, 30, 47, 48, 60, 61).

Every format is generated, watermarked, re-opened with its native library and
checked for structure preservation, diagonal + semi-transparent marking, the
correct identity fields (short wallet only), a decodable QR, and an unchanged original.
"""

import hashlib
import io
import re
import zipfile
from datetime import datetime, UTC

import cv2
import numpy as np
import pytest

from app.services.watermark import (
    HistoryEntry, UnsupportedDocumentType, WatermarkContext, render_protected_copy, short_wallet,
)

VIEWER_WALLET = "0xb8b8cf51335fa22e7967a9e72c16f738ddc7ef38"
OWNER_WALLET = "0xa72f00000000000000000000000000000000091cd"
MOD_WALLET_1 = "0xa72f11111111111111111111111111111111191cd"
MOD_WALLET_2 = "0xf39f22222222222222222222222222222222ab12"
WHEN = datetime(2026, 9, 22, 10, 30, tzinfo=UTC)
VERIFY_URL = "http://localhost:5173/verify/action/VW-20260922-0001"


def ctx(kind="VIEW", history=None, action_id="VW-20260922-0001"):
    return WatermarkContext(
        kind=kind, actor_name="Priya Sharma", actor_wallet=VIEWER_WALLET, client_ip="203.0.113.44",
        action_id=action_id, verify_url=VERIFY_URL.replace("VW-20260922-0001", action_id), when=WHEN,
        owner_name="Owner Name" if kind == "DOWNLOAD" else None,
        owner_wallet=OWNER_WALLET if kind == "DOWNLOAD" else None,
        history=history or [],
    )


def sha(b):
    return hashlib.sha256(b).hexdigest()


def decode_qrs_from_image_bytes(img_bytes):
    arr = cv2.imdecode(np.frombuffer(img_bytes, np.uint8), cv2.IMREAD_COLOR)
    det = cv2.QRCodeDetector()
    found = [det.detectAndDecode(arr)[0]]
    ok, multi, _, _ = det.detectAndDecodeMulti(arr)
    found += list(multi) if ok else []
    return [d for d in found if d]


def assert_identity(text: str, c: WatermarkContext):
    assert "CONFIDENTIAL" in text
    assert c.actor_name in text
    assert short_wallet(c.actor_wallet) in text
    assert c.actor_wallet not in text and c.actor_wallet[2:] not in text  # never the full wallet
    assert c.client_ip in text
    assert "22 September 2026" in text
    assert c.action_id in text


# ------------------------------------------------------------------ fixtures: sample documents

def sample_pdf(pages=3):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    for i in range(pages):
        c.drawString(72, 750, f"Original page {i + 1} body text")
        c.showPage()
    c.save()
    return buf.getvalue()


def sample_docx():
    import docx
    d = docx.Document()
    d.add_heading("Radar Specification", 0)
    for i in range(3):
        d.add_paragraph(f"Original paragraph {i + 1}")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "cell-a"
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


def sample_pptx(slides=3):
    from pptx import Presentation
    from pptx.util import Inches
    p = Presentation()
    for i in range(slides):
        s = p.slides.add_slide(p.slide_layouts[1])
        s.shapes.title.text = f"Slide {i + 1} title"
        s.shapes.add_textbox(Inches(1), Inches(3), Inches(4), Inches(1)).text_frame.text = f"Original body {i + 1}"
    out = io.BytesIO()
    p.save(out)
    return out.getvalue()


def sample_xlsx():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws["A1"], ws["A2"], ws["A3"] = 10, 32, "=A1+A2"
    wb.create_sheet("Notes")["B2"] = "keep me"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def sample_image(fmt):
    from PIL import Image
    img = Image.new("RGB", (900, 600), (240, 248, 255))
    out = io.BytesIO()
    img.save(out, format=fmt)
    return out.getvalue()


# ------------------------------------------------------------------ tests

def test_pdf_every_page_diagonal_translucent_qr_and_control_page():
    from pypdf import PdfReader
    src = sample_pdf(3)
    before = sha(src)
    c = ctx()
    out, mime, name = render_protected_copy(src, "spec.pdf", c)
    assert sha(src) == before and out != src
    assert mime == "application/pdf" and name.endswith(".pdf")
    r = PdfReader(io.BytesIO(out))
    assert len(r.pages) == 4  # 3 originals + control page
    for i, page in enumerate(r.pages[:3]):
        text = page.extract_text()
        assert f"Original page {i + 1} body text" in text  # content preserved
        assert_identity(text, c)
        content = page.get_contents().get_data()
        assert re.search(rb"0?\.7071\d* 0?\.7071\d* -0?\.7071\d* 0?\.7071\d* [\d.]+ [\d.]+ cm", content), "no 45 deg rotation"
        gs = page["/Resources"]["/ExtGState"]
        alphas = [float(v.get_object().get("/ca", 1)) for v in gs.values()]
        assert any(0 < a < 1 for a in alphas), "no semi-transparent fill"
        xobjs = page["/Resources"]["/XObject"]
        assert any(x.get_object()["/Subtype"] == "/Image" for x in xobjs.values()), "no QR image"
    assert "LexCorp Document Control" in r.pages[3].extract_text()
    # QR decodes to the verification reference
    img = next(x.get_object() for x in r.pages[0]["/Resources"]["/XObject"].values() if x.get_object()["/Subtype"] == "/Image")
    from pypdf.generic import IndirectObject  # noqa
    pil = r.pages[0].images[0].image
    buf = io.BytesIO(); pil.convert("RGB").resize((400, 400)).save(buf, format="PNG")
    assert VERIFY_URL in decode_qrs_from_image_bytes(buf.getvalue())


def test_view_watermark_never_shows_owner_but_download_does():
    from pypdf import PdfReader
    v, _, _ = render_protected_copy(sample_pdf(1), "a.pdf", ctx("VIEW"))
    d, _, _ = render_protected_copy(sample_pdf(1), "a.pdf", ctx("DOWNLOAD", action_id="DL-20260922-0001"))
    vt = "".join(p.extract_text() for p in PdfReader(io.BytesIO(v)).pages)
    dt = "".join(p.extract_text() for p in PdfReader(io.BytesIO(d)).pages)
    assert "Viewed by: Priya Sharma" in vt and "View ID: VW-20260922-0001" in vt
    assert "Owner" not in vt and short_wallet(OWNER_WALLET) not in vt
    assert "Downloaded by: Priya Sharma" in dt and "Download ID: DL-20260922-0001" in dt
    assert f"Owner: Owner Name ({short_wallet(OWNER_WALLET)})" in dt


def test_text_rendered_as_protected_pdf():
    from pypdf import PdfReader
    src = b"line one of the secret\nline two\n"
    c = ctx()
    out, mime, name = render_protected_copy(src, "notes.txt", c)
    assert mime == "application/pdf" and name.endswith(".pdf")
    r = PdfReader(io.BytesIO(out))
    text = r.pages[0].extract_text()
    assert "line one of the secret" in text
    assert_identity(text, c)


def test_docx_structure_preserved_with_native_watermark_and_control_section():
    import docx
    src = sample_docx()
    before = sha(src)
    c = ctx()
    out, mime, _ = render_protected_copy(src, "spec.docx", c)
    assert sha(src) == before
    assert mime.endswith("wordprocessingml.document")
    assert zipfile.ZipFile(io.BytesIO(out)).testzip() is None
    d = docx.Document(io.BytesIO(out))  # reopens
    texts = [p.text for p in d.paragraphs]
    assert texts[:4] == ["Radar Specification", "Original paragraph 1", "Original paragraph 2", "Original paragraph 3"]
    assert d.tables[0].cell(0, 0).text == "cell-a"
    header_xml = d.sections[0].header._element.xml
    assert "rotation:315" in header_xml and 'opacity=".25"' in header_xml
    assert c.action_id in header_xml and "LexCorpWatermark" in header_xml
    body = "\n".join(texts)
    assert "LexCorp Document Control" in body
    assert_identity(body, c)
    assert any(r.reltype.endswith("/image") for r in d.sections[0].header.part.rels.values())  # QR in header


def test_pptx_every_slide_marked_structure_preserved():
    from pptx import Presentation
    src = sample_pptx(3)
    c = ctx()
    out, mime, _ = render_protected_copy(src, "deck.pptx", c)
    p = Presentation(io.BytesIO(out))
    slides = list(p.slides)
    assert len(slides) == 4
    for i, s in enumerate(slides[:3]):
        names = [sh.name for sh in s.shapes]
        assert s.shapes.title.text == f"Slide {i + 1} title"
        assert any(sh.has_text_frame and sh.text_frame.text == f"Original body {i + 1}" for sh in s.shapes)
        wm = next(sh for sh in s.shapes if sh.name == "LexCorpWatermark")
        assert wm.rotation == 315
        assert "<a:alpha" in wm._element.xml
        assert c.action_id in wm.text_frame.text and short_wallet(VIEWER_WALLET) in wm.text_frame.text
        assert "LexCorpWatermarkQR" in names
    control_text = "\n".join(sh.text_frame.text for sh in slides[3].shapes if sh.has_text_frame)
    assert_identity(control_text, c)


def test_xlsx_values_formulas_preserved_header_and_control_sheet():
    import openpyxl
    src = sample_xlsx()
    c = ctx()
    out, mime, _ = render_protected_copy(src, "budget.xlsx", c)
    wb = openpyxl.load_workbook(io.BytesIO(out))
    assert wb.sheetnames == ["Budget", "Notes", "LexCorp Control"]
    b = wb["Budget"]
    assert b["A1"].value == 10 and b["A3"].value == "=A1+A2" and wb["Notes"]["B2"].value == "keep me"
    assert c.action_id in b.oddHeader.center.text and short_wallet(VIEWER_WALLET) in b.oddHeader.center.text
    control = "\n".join(str(r[0].value) for r in wb["LexCorp Control"].iter_rows() if r[0].value)
    assert_identity(control, c)
    assert any(n.startswith("xl/media/") for n in zipfile.ZipFile(io.BytesIO(out)).namelist())  # QR image


@pytest.mark.parametrize("fmt,mime", [("PNG", "image/png"), ("JPEG", "image/jpeg")])
def test_image_overlay_band_and_qr(fmt, mime):
    from PIL import Image
    src = sample_image(fmt)
    c = ctx()
    out, got_mime, _ = render_protected_copy(src, f"photo.{fmt.lower()}", c)
    assert got_mime == mime
    img = Image.open(io.BytesIO(out)).convert("RGB")
    orig = Image.open(io.BytesIO(src)).convert("RGB")
    assert img.width == orig.width and img.height > orig.height  # footer band added, content not covered
    region = np.asarray(img)[: orig.height].astype(int)
    base = np.asarray(orig).astype(int)
    diff = np.abs(region - base).mean()
    assert 1 < diff < 60, diff  # visible but translucent (content still visible underneath)
    assert VERIFY_URL in decode_qrs_from_image_bytes(out)


def test_modification_history_cumulative_non_overlapping_and_deterministic():
    from pypdf import PdfReader
    h1 = HistoryEntry("MOD-20260922-0001", MOD_WALLET_1, datetime(2026, 9, 22, 9, tzinfo=UTC))
    h2 = HistoryEntry("MOD-20260923-0002", MOD_WALLET_2, datetime(2026, 9, 23, 9, tzinfo=UTC))
    # Revision 1 copy
    c1 = ctx("DOWNLOAD", history=[h1], action_id="DL-20260922-0001")
    # Revision 2 copy (history given out of order: output must still be oldest-first)
    c2 = ctx("DOWNLOAD", history=[h2, h1], action_id="DL-20260923-0002")
    assert c2.all_lines() == ctx("DOWNLOAD", history=[h1, h2], action_id="DL-20260923-0002").all_lines()
    t1 = "".join(p.extract_text() for p in PdfReader(io.BytesIO(render_protected_copy(sample_pdf(1), "a.pdf", c1)[0])).pages)
    t2 = "".join(p.extract_text() for p in PdfReader(io.BytesIO(render_protected_copy(sample_pdf(1), "a.pdf", c2)[0])).pages)
    assert "Modification ID: MOD-20260922-0001" in t1 and "MOD-20260923-0002" not in t1
    i1, i2 = t2.index("MOD-20260922-0001"), t2.index("MOD-20260923-0002")
    assert i1 < i2  # previous history retained, in order
    assert f"Modified by wallet: {short_wallet(MOD_WALLET_1)}" in t2 and "Date: 23-09-2026" in t2
    assert MOD_WALLET_2 not in t2


def test_unsupported_types_refused_not_passed_through():
    for data, name in [(b"MZ\x90\x00binary", "tool.exe"), (b"\xd0\xcf\x11\xe0legacy", "old.doc")]:
        with pytest.raises(UnsupportedDocumentType):
            render_protected_copy(data, name, ctx())
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("a.txt", "x")
    with pytest.raises(UnsupportedDocumentType):
        render_protected_copy(buf.getvalue(), "bundle.zip", ctx())



def test_xlsx_charts_images_and_parts_preserved_byte_for_byte():
    """The XLSX watermark no longer round-trips through openpyxl: charts, images and
    every original part survive unchanged; only sheet XML gains header/footer."""
    import openpyxl
    from openpyxl.chart import BarChart, Reference
    from openpyxl.drawing.image import Image as XLImage
    from PIL import Image
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    for i in range(1, 6):
        ws.append([f"Item {i}", i * 10])
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=5))
    ws.add_chart(chart, "D2")
    pic = io.BytesIO()
    Image.new("RGB", (40, 40), (10, 60, 120)).save(pic, format="PNG")
    ws.add_image(XLImage(io.BytesIO(pic.getvalue())), "D20")
    src_buf = io.BytesIO()
    wb.save(src_buf)
    src = src_buf.getvalue()
    before = sha(src)
    c = ctx()
    out, mime, _ = render_protected_copy(src, "report.xlsx", c)
    assert sha(src) == before
    zs, zo = zipfile.ZipFile(io.BytesIO(src)), zipfile.ZipFile(io.BytesIO(out))
    for name in zs.namelist():
        assert name in zo.namelist(), name
        if name.startswith(("xl/charts/", "xl/media/", "xl/drawings/", "xl/styles", "xl/sharedStrings", "xl/theme/")):
            assert zs.read(name) == zo.read(name), f"{name} changed"
    sheet_xml = zo.read("xl/worksheets/sheet1.xml").decode()
    assert c.action_id in sheet_xml and "<headerFooter>" in sheet_xml.replace(' xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"', "")
    # headerFooter sits before <drawing> (schema order) and the workbook still opens with its chart
    assert sheet_xml.index("headerFooter") < sheet_xml.index("<drawing")
    wb2 = openpyxl.load_workbook(io.BytesIO(out))
    assert wb2.sheetnames == ["Data", "LexCorp Control"]
    assert wb2["Data"]["B5"].value == 50
    assert len(wb2["Data"]._charts) == 1
    control = "\n".join(str(r[0].value) for r in wb2["LexCorp Control"].iter_rows() if r[0].value)
    assert_identity(control, c)
