"""Derived, watermarked representations of vault documents.

The watermark is a DERIVED security/traceability representation. Its content comes
only from authoritative server-side records (the acting identity, the effective
client IP from the trusted request context, backend-generated action IDs and the
modification history stored in the database) - never from text inside an uploaded
file. The original bytes are never modified; every call returns a new copy.

Per format (all write a diagonal, semi-transparent marking and a QR code that
carries only a verification reference):

  PDF   every page: diagonal semi-transparent text + QR in the bottom-right corner
        + compact identity line; an appended "Document Control" page holds the full
        block and the cumulative modification history.
  DOCX  every section header: a native Word watermark shape (VML text path,
        rotated 315 deg, fill opacity 0.25) + header line with QR; an appended
        control section with the full block and history.
  PPTX  every slide: rotated semi-transparent text box + QR; an appended control
        slide. Existing slides and shapes are left in place.
  XLSX  edited at the OOXML package level, so charts, pivot tables, images and every
        other part are preserved byte-for-byte: each sheet's header/footer carries
        the watermark (visible in Page Layout / print) and an appended
        "LexCorp Control" sheet holds the block, history and QR. Cells are not
        overlaid so the workbook stays usable.
  Image diagonal tiled semi-transparent text over the image; the block and QR go in
        an added footer band so the image content itself is not covered.
  Text  rendered to a PDF ("protected rendered representation") and watermarked.

Unsupported formats raise UnsupportedDocumentType; the caller must refuse rather
than return an unwatermarked original.

Removable, verifiable layers. Every element added is a separate, labelled layer tagged
with the action ID (PDF: a Form XObject per page drawn from its own content stream, and
tagged control pages; Word: content controls; PowerPoint: named shapes and a named
control slide; Excel: the headers/footers set and a named control sheet; images: a
regenerable overlay and a footer band). render_protected_copy_with_manifest() also
returns a MANIFEST: fingerprints of every element as issued. When a modified copy comes
back, app.services.watermark_strip verifies each element against the manifest (reporting
where anything was changed or removed) and removes the layers, so a revision is stored
clean and later copies carry one watermark, never a stack.
"""

from __future__ import annotations


import hashlib
import io
import logging
import textwrap
import zipfile
from dataclasses import dataclass, field
from datetime import datetime, UTC
from typing import List, Optional

from app.services.qr import generate_qr_for_url

logger = logging.getLogger(__name__)

WATERMARK_TAG = "LexCorpWatermark"
CONTROL_TITLE = "LexCorp Document Control"
LAYER_KEY = "/LexCorpWM"          # PDF: marks our content streams and overlay forms
CONTROL_KEY = "/LexCorpControl"   # PDF: marks our control pages


def sha256_hex(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()
_MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August",
           "September", "October", "November", "December"]


class UnsupportedDocumentType(Exception):
    """The file type has no safe watermarking implementation."""


def short_wallet(address: Optional[str]) -> str:
    """0xb8b8cf5...c7b5df -> 0xb8b8...b5df (never expose the full wallet)."""
    if not address:
        return "unknown"
    a = address.strip()
    return f"{a[:6]}...{a[-4:]}" if len(a) > 12 else a


def long_date(dt: datetime) -> str:
    return f"{dt.day:02d} {_MONTHS[dt.month - 1]} {dt.year}"


def short_date(dt: datetime) -> str:
    return dt.strftime("%d-%m-%Y")


@dataclass
class HistoryEntry:
    modification_id: str
    modifier_wallet: str
    when: datetime
    modifier_name: Optional[str] = None
    version: Optional[int] = None


@dataclass
class WatermarkContext:
    kind: str                      # "VIEW" | "DOWNLOAD"
    actor_name: str
    actor_wallet: str
    client_ip: str
    action_id: str
    verify_url: str
    when: datetime = field(default_factory=lambda: datetime.now(UTC))
    owner_name: Optional[str] = None     # DOWNLOAD only: never shown on views
    owner_wallet: Optional[str] = None
    history: List[HistoryEntry] = field(default_factory=list)

    def primary_lines(self) -> List[str]:
        verb, label = ("Viewed by", "View ID") if self.kind == "VIEW" else ("Downloaded by", "Download ID")
        lines = [
            "CONFIDENTIAL",
            f"{verb}: {self.actor_name}",
            f"Wallet: {short_wallet(self.actor_wallet)}",
            f"IP: {self.client_ip}",
            f"Date: {long_date(self.when)}",
            f"{label}: {self.action_id}",
        ]
        if self.kind == "DOWNLOAD" and self.owner_wallet:
            lines.append(f"Owner: {self.owner_name or 'Unnamed'} ({short_wallet(self.owner_wallet)})")
        return lines

    def history_blocks(self) -> List[List[str]]:
        """Cumulative, deterministic (oldest first) modification history."""
        ordered = sorted(self.history, key=lambda h: (h.when, h.modification_id))
        return [[
            f"Modification ID: {h.modification_id}",
            f"Modified by wallet: {short_wallet(h.modifier_wallet)}",
            f"Date: {short_date(h.when)}",
        ] for h in ordered]

    def all_lines(self) -> List[str]:
        out = list(self.primary_lines())
        for block in self.history_blocks():
            out.append("")
            out.extend(block)
        return out

    def compact(self) -> str:
        verb = "Viewed by" if self.kind == "VIEW" else "Downloaded by"
        return (f"CONFIDENTIAL | {verb} {self.actor_name} | {short_wallet(self.actor_wallet)} | "
                f"IP {self.client_ip} | {long_date(self.when)} | {self.action_id}")

    def qr_png(self) -> bytes:
        return generate_qr_for_url(self.verify_url)


# ------------------------------------------------------------------ type detection

def detect_type(data: bytes, filename: str) -> str:
    name = (filename or "").lower()
    if data[:5] == b"%PDF-":
        return "pdf"
    if data[:8] == b"\x89PNG\r\n\x1a\n" or data[:3] == b"\xff\xd8\xff":
        return "image"
    if data[:4] == b"PK\x03\x04":
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                names = set(z.namelist())
        except zipfile.BadZipFile:
            raise UnsupportedDocumentType("Corrupt ZIP container")
        if "word/document.xml" in names:
            return "docx"
        if "ppt/presentation.xml" in names:
            return "pptx"
        if "xl/workbook.xml" in names:
            return "xlsx"
        raise UnsupportedDocumentType("ZIP archives other than DOCX/PPTX/XLSX are not supported")
    if name.endswith((".gif", ".bmp", ".webp", ".tif", ".tiff")):
        return "image"
    if name.endswith((".txt", ".md", ".csv", ".log", ".json", ".xml")):
        try:
            data.decode("utf-8")
            return "text"
        except UnicodeDecodeError:
            raise UnsupportedDocumentType("Text file is not valid UTF-8")
    raise UnsupportedDocumentType(f"No safe watermarking implementation for '{filename}'")


# ------------------------------------------------------------------ PDF

def _pdf_overlay(width: float, height: float, ctx: WatermarkContext, qr_png: bytes) -> bytes:
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=(width, height))
    # Diagonal, semi-transparent marking across the page
    c.saveState()
    c.setFillColorRGB(0.75, 0.05, 0.05)
    c.setFillAlpha(0.16)
    c.translate(width / 2, height / 2)
    c.rotate(45)
    size = max(18, min(width, height) / 9)
    c.setFont("Helvetica-Bold", size)
    c.drawCentredString(0, size * 0.6, "CONFIDENTIAL")
    c.setFont("Helvetica", size * 0.32)
    for i, line in enumerate(ctx.primary_lines()[1:]):
        c.drawCentredString(0, -size * 0.2 - i * size * 0.42, line)
    c.restoreState()
    # Compact identity line and QR in the bottom margin
    qr = max(36, min(width, height) * 0.075)
    c.saveState()
    c.setFillAlpha(0.7)
    c.setFillColorRGB(0.45, 0.05, 0.05)
    c.setFont("Helvetica", 6.5)
    c.drawString(8, 8, ctx.compact()[: int((width - qr - 24) / 3.2)])
    c.restoreState()
    c.drawImage(ImageReader(io.BytesIO(qr_png)), width - qr - 6, 6, qr, qr)
    c.showPage()
    c.save()
    return buf.getvalue()


def _pdf_control_pages(width: float, height: float, ctx: WatermarkContext, qr_png: bytes) -> bytes:
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=(width, height))
    y = height - 50
    c.setFont("Helvetica-Bold", 14)
    c.drawString(40, y, CONTROL_TITLE)
    y -= 12
    c.setFont("Helvetica", 8)
    c.drawString(40, y, "Derived traceability copy. Authoritative record: LexCorp database, audit trail and blockchain references.")
    y -= 24
    qr = 110
    c.drawImage(ImageReader(io.BytesIO(qr_png)), width - qr - 40, height - qr - 40, qr, qr)
    for line in ctx.all_lines():
        if y < 50:
            c.showPage()
            y = height - 50
        c.setFont("Helvetica-Bold" if line in ("CONFIDENTIAL",) else "Helvetica", 11 if line == "CONFIDENTIAL" else 9.5)
        c.drawString(40, y, line)
        y -= 14 if line else 8
    c.showPage()
    c.save()
    return buf.getvalue()


def _pdf_stream(data: bytes, tag):
    from pypdf.generic import DecodedStreamObject, NameObject
    st = DecodedStreamObject()
    st.set_data(data)
    st[NameObject(LAYER_KEY)] = tag
    return st


def watermark_pdf(data: bytes, ctx: WatermarkContext, manifest: Optional[dict] = None) -> bytes:
    """Each page gets its overlay as a separate Form XObject (/LCWM<n>) drawn by its own
    tagged content stream; the page's original content is wrapped in q/Q and otherwise
    untouched. Control pages are appended and tagged. Everything can be found and removed."""
    from pypdf import PdfReader, PdfWriter
    from pypdf.generic import ArrayObject, DictionaryObject, FloatObject, NameObject, TextStringObject

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            raise UnsupportedDocumentType("Encrypted PDF cannot be watermarked")
    qr_png = ctx.qr_png()
    tag = TextStringObject(ctx.action_id)
    writer = PdfWriter(clone_from=reader)
    overlays = {}
    for n, page in enumerate(writer.pages, start=1):
        box = page.mediabox
        w, h, x0, y0 = float(box.width), float(box.height), float(box.left), float(box.bottom)
        ov = PdfReader(io.BytesIO(_pdf_overlay(w, h, ctx, qr_png))).pages[0]
        form = _pdf_stream(ov.get_contents().get_data(), tag)
        form.update({
            NameObject("/Type"): NameObject("/XObject"), NameObject("/Subtype"): NameObject("/Form"),
            NameObject("/BBox"): ArrayObject([FloatObject(0), FloatObject(0), FloatObject(w), FloatObject(h)]),
            NameObject("/Resources"): ov["/Resources"].clone(writer),
        })
        name = f"LCWM{n}"
        form_ref = writer._add_object(form)
        if "/Resources" not in page:
            page[NameObject("/Resources")] = DictionaryObject()
        resources = page["/Resources"].get_object()
        if "/XObject" not in resources:
            resources[NameObject("/XObject")] = DictionaryObject()
        resources["/XObject"].get_object()[NameObject("/" + name)] = form_ref
        original = page.get("/Contents")
        original = [] if original is None else (list(original.get_object()) if isinstance(original.get_object(), ArrayObject) else [original])
        opener = writer._add_object(_pdf_stream(b"q\n", tag))
        closer = writer._add_object(_pdf_stream(f"\nQ\nq 1 0 0 1 {x0:g} {y0:g} cm /{name} Do Q\n".encode(), tag))
        page[NameObject("/Contents")] = ArrayObject([opener, *original, closer])
        images = form["/Resources"].get_object().get("/XObject", {})
        overlays[name] = {"page": n, "form": sha256_hex(form.get_data()),
                          "images": sorted(sha256_hex(images[k].get_object().get_data()) for k in images)}

    first = reader.pages[0].mediabox if reader.pages else None
    cw, ch = (float(first.width), float(first.height)) if first else (595.0, 842.0)
    controls = []
    for cp in PdfReader(io.BytesIO(_pdf_control_pages(cw, ch, ctx, qr_png))).pages:
        added = writer.add_page(cp)
        added[NameObject(CONTROL_KEY)] = tag
        controls.append(sha256_hex(added.get_contents().get_data()))
    writer.add_metadata({"/Subject": f"LexCorp {ctx.kind} copy {ctx.action_id}", "/Keywords": WATERMARK_TAG})
    if manifest is not None:
        manifest.update({"format": "pdf", "page_count": len(reader.pages), "overlays": overlays, "control_pages": controls})
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


# ------------------------------------------------------------------ Text -> PDF

def text_to_pdf(data: bytes, title: str) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    text = data.decode("utf-8", errors="replace").replace("\t", "    ")
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    w, h = A4
    c.setTitle(title)
    y = h - 60
    c.setFont("Courier", 9)
    for raw in text.splitlines() or [""]:
        for line in (textwrap.wrap(raw, 95, replace_whitespace=False, drop_whitespace=False) or [""]):
            if y < 60:
                c.showPage()
                c.setFont("Courier", 9)
                y = h - 60
            c.drawString(40, y, line)
            y -= 11
    c.showPage()
    c.save()
    return buf.getvalue()


# ------------------------------------------------------------------ Images

def _font(size: int):
    from PIL import ImageFont
    for name in ("arial.ttf", "DejaVuSans.ttf", "LiberationSans-Regular.ttf"):
        try:
            return ImageFont.truetype(name, size)
        except OSError:
            continue
    return ImageFont.load_default(size=size)


def image_overlay_layer(w: int, h: int, text: str, size: int):
    """The diagonal tiled marking for a w x h image (deterministic, so it can be regenerated
    exactly to verify and remove it)."""
    from PIL import Image, ImageDraw
    diag = int((w ** 2 + h ** 2) ** 0.5)
    layer = Image.new("RGBA", (diag, diag), (0, 0, 0, 0))
    d = ImageDraw.Draw(layer)
    font = _font(size)
    step = int(size * 3.2)
    for yy in range(0, diag, step):
        d.text((0 if (yy // step) % 2 else step, yy), (text + "     ") * 6, font=font, fill=(200, 0, 0, 64))
    layer = layer.rotate(45, expand=False)
    left, top = (diag - w) // 2, (diag - h) // 2
    return layer.crop((left, top, left + w, top + h))


def image_band(w: int, lines: List[str], qr_png: bytes):
    """The footer band (block text + QR) added below the image (deterministic)."""
    from PIL import Image, ImageDraw
    fsize = max(11, min(w, 1600) // 70)
    ffont = _font(fsize)
    qr = Image.open(io.BytesIO(qr_png)).convert("RGBA")
    band_h = max(len(lines) * (fsize + 4) + 20, int(fsize * 9))
    qr_side = min(band_h - 16, max(96, band_h - 16))
    qr = qr.resize((qr_side, qr_side))
    band = Image.new("RGBA", (w, band_h), (253, 242, 242, 255))
    dd = ImageDraw.Draw(band)
    y = 10
    for line in lines:
        dd.text((12, y), line, font=ffont, fill=(127, 29, 29, 255))
        y += fsize + 4
    band.paste(qr, (max(0, w - qr_side - 8), 8), qr)
    return band


def watermark_image(data: bytes, ctx: WatermarkContext, manifest: Optional[dict] = None) -> tuple[bytes, str]:
    from PIL import Image

    src = Image.open(io.BytesIO(data))
    fmt = (src.format or "PNG").upper()
    base = src.convert("RGBA")
    w, h = base.size
    size = max(14, min(w, h) // 18)
    text = f"CONFIDENTIAL  {ctx.actor_name}  {short_wallet(ctx.actor_wallet)}  {ctx.action_id}"
    marked = Image.alpha_composite(base, image_overlay_layer(w, h, text, size))
    lines = ctx.all_lines()
    band = image_band(w, lines, ctx.qr_png())
    canvas = Image.new("RGBA", (w, h + band.height), (255, 255, 255, 255))
    canvas.paste(marked, (0, 0))
    canvas.paste(band, (0, h))
    out = io.BytesIO()
    jpeg = fmt in ("JPEG", "JPG")
    if jpeg:
        canvas.convert("RGB").save(out, format="JPEG", quality=92)
    else:
        canvas.save(out, format="PNG")
    if manifest is not None:
        manifest.update({"format": "image", "width": w, "height": h, "band_height": band.height, "jpeg": jpeg,
                         "overlay_text": text, "overlay_size": size, "band_lines": lines})
    return out.getvalue(), ("image/jpeg" if jpeg else "image/png")


# ------------------------------------------------------------------ DOCX

_VML_NS = "urn:schemas-microsoft-com:vml"
_O_NS = "urn:schemas-microsoft-com:office:office"


def _docx_watermark_run(text: str):
    from docx.oxml import parse_xml

    safe = (text.replace("&", "&amp;").replace('"', "&quot;").replace("<", "&lt;").replace(">", "&gt;"))
    xml = (
        '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        f'xmlns:v="{_VML_NS}" xmlns:o="{_O_NS}"><w:pict>'
        f'<v:shape id="{WATERMARK_TAG}" o:spid="_x0000_s4097" type="#_x0000_t136" '
        'style="position:absolute;margin-left:0;margin-top:0;width:468pt;height:117pt;rotation:315;'
        'z-index:-251654144;mso-position-horizontal:center;mso-position-horizontal-relative:margin;'
        'mso-position-vertical:center;mso-position-vertical-relative:margin" o:allowincell="f" '
        'fillcolor="#c00000" stroked="f"><v:fill opacity=".25"/>'
        f'<v:textpath style="font-family:&quot;Calibri&quot;;font-size:1pt" string="{safe}"/>'
        '</v:shape></w:pict></w:r>'
    )
    return parse_xml(xml)


_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def docx_sdt(tag: str, alias: str):
    """An empty block-level content control carrying our tag (Word keeps tags on save)."""
    from docx.oxml import parse_xml
    return parse_xml(f'<w:sdt xmlns:w="{_W_NS}"><w:sdtPr><w:alias w:val="{alias}"/><w:tag w:val="{tag}"/>'
                     '<w:lock w:val="sdtContentLocked"/></w:sdtPr><w:sdtContent/></w:sdt>')


def _into_sdt(elements, tag: str, alias: str):
    sdt = docx_sdt(tag, alias)
    elements[0].addprevious(sdt)
    content = sdt[-1]
    for el in elements:
        content.append(el)
    return sdt


def watermark_docx(data: bytes, ctx: WatermarkContext, manifest: Optional[dict] = None) -> bytes:
    import docx
    from docx.enum.text import WD_BREAK
    from docx.shared import Inches, Pt, RGBColor

    doc = docx.Document(io.BytesIO(data))
    qr_png = ctx.qr_png()
    shape_text = f"CONFIDENTIAL {ctx.actor_name} {ctx.action_id}"
    seen = set()
    for section in doc.sections:
        for header in (section.header, section.first_page_header, section.even_page_header):
            part_id = id(header._element)
            if part_id in seen:
                continue
            seen.add(part_id)
            header.is_linked_to_previous = False
            p = header.add_paragraph()
            p._p.append(_docx_watermark_run(shape_text))
            run = p.add_run(ctx.compact() + "  ")
            run.font.size = Pt(7)
            run.font.color.rgb = RGBColor(0x99, 0x1B, 0x1B)
            p.add_run().add_picture(io.BytesIO(qr_png), width=Inches(0.55))
            _into_sdt([p._p], f"LexCorpWM:{ctx.action_id}", "LexCorp watermark")

    # Dedicated control section: full block + cumulative history (never overlays content)
    body = doc.element.body
    before = len(body)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading(CONTROL_TITLE, level=1)
    for line in ctx.all_lines():
        para = doc.add_paragraph()
        r = para.add_run(line)
        r.bold = line == "CONFIDENTIAL"
    doc.add_picture(io.BytesIO(qr_png), width=Inches(1.4))
    added = [el for el in list(body)[before - 1:] if el.tag != f"{{{_W_NS}}}sectPr"][-(len(body) - before):]
    _into_sdt(added, f"LexCorpControl:{ctx.action_id}", "LexCorp document control")
    if manifest is not None:
        manifest.update({"format": "docx", "header_text": (ctx.compact() + "  ").strip(), "shape_string": shape_text,
                         "headers": len(seen), "control_lines": [CONTROL_TITLE] + [x for x in ctx.all_lines() if x],
                         "qr": sha256_hex(qr_png)})
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


# ------------------------------------------------------------------ PPTX

def _set_alpha(run_or_shape_fill_xml_parent, alpha_pct: int):
    """Add <a:alpha val=.../> to the first srgbClr under the element."""
    from pptx.oxml.ns import qn
    for clr in run_or_shape_fill_xml_parent.iter(qn("a:srgbClr")):
        alpha = clr.makeelement(qn("a:alpha"), {"val": str(alpha_pct * 1000)})
        clr.append(alpha)
        return


def watermark_pptx(data: bytes, ctx: WatermarkContext, manifest: Optional[dict] = None) -> bytes:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Emu, Pt

    prs = Presentation(io.BytesIO(data))
    sw, sh = prs.slide_width, prs.slide_height
    qr_png = ctx.qr_png()
    text = f"CONFIDENTIAL · {ctx.actor_name} · {short_wallet(ctx.actor_wallet)} · {ctx.action_id}"
    original_slides = list(prs.slides)
    for n, slide in enumerate(original_slides, start=1):
        box = slide.shapes.add_textbox(Emu(int(sw * 0.05)), Emu(int(sh * 0.4)), Emu(int(sw * 0.9)), Emu(int(sh * 0.2)))
        box.name = f"LexCorpWM:{ctx.action_id}:{n}"
        box.rotation = 315
        tf = box.text_frame
        tf.word_wrap = True
        run = tf.paragraphs[0].add_run()
        run.text = text
        run.font.size = Pt(max(14, int(sh / 12700 / 16)))
        run.font.bold = True
        run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
        _set_alpha(run._r, 25)
        qr_side = Emu(int(min(sw, sh) * 0.1))
        pic = slide.shapes.add_picture(io.BytesIO(qr_png), sw - qr_side - Emu(60000), sh - qr_side - Emu(60000), qr_side, qr_side)
        pic.name = f"LexCorpWMQR:{ctx.action_id}:{n}"

    # Appended control slide (original slides untouched otherwise)
    layout = min(prs.slide_layouts, key=lambda l: len(l.placeholders))
    control = prs.slides.add_slide(layout)
    for ph in list(control.placeholders):
        ph._element.getparent().remove(ph._element)
    tb = control.shapes.add_textbox(Emu(int(sw * 0.05)), Emu(int(sh * 0.05)), Emu(int(sw * 0.62)), Emu(int(sh * 0.9)))
    tb.name = f"LexCorpControl:{ctx.action_id}"
    tf = tb.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].text = CONTROL_TITLE
    tf.paragraphs[0].runs[0].font.bold = True
    for line in ctx.all_lines():
        p = tf.add_paragraph()
        p.text = line
        if p.runs:
            p.runs[0].font.size = Pt(12)
    side = Emu(int(min(sw, sh) * 0.3))
    cqr = control.shapes.add_picture(io.BytesIO(qr_png), sw - side - Emu(int(sw * 0.05)), Emu(int(sh * 0.08)), side, side)
    cqr.name = f"LexCorpControlQR:{ctx.action_id}"
    if manifest is not None:
        manifest.update({"format": "pptx", "slide_count": len(original_slides), "box_text": text,
                         "control_lines": [CONTROL_TITLE] + [x for x in ctx.all_lines() if x], "qr": sha256_hex(qr_png)})
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


# ------------------------------------------------------------------ XLSX

_NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_NS_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_NS_PKG_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
_NS_CT = "http://schemas.openxmlformats.org/package/2006/content-types"
_NS_XDR = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
_NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
_REL_WS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet"
_REL_DRAWING = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing"
_REL_IMAGE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
# CT_Worksheet / CT_Chartsheet children that must come AFTER <headerFooter>.
_AFTER_HF = {"rowBreaks", "colBreaks", "customProperties", "cellWatches", "ignoredErrors", "smartTags", "drawing",
             "legacyDrawing", "legacyDrawingHF", "drawingHF", "picture", "oleObjects", "controls",
             "webPublishItems", "tableParts", "extLst"}


def _hf_escape(text: str) -> str:
    return text.replace("&", "&&")


def _xlsx_set_header_footer(xml: bytes, header: str, footer: str) -> bytes:
    from lxml import etree
    root = etree.fromstring(xml)
    for old in root.findall(f"{{{_NS_MAIN}}}headerFooter"):
        root.remove(old)
    hf = etree.Element(f"{{{_NS_MAIN}}}headerFooter")
    etree.SubElement(hf, f"{{{_NS_MAIN}}}oddHeader").text = header
    etree.SubElement(hf, f"{{{_NS_MAIN}}}oddFooter").text = footer
    anchor = next((c for c in root if isinstance(c.tag, str) and etree.QName(c).localname in _AFTER_HF), None)
    if anchor is not None:
        anchor.addprevious(hf)
    else:
        root.append(hf)
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def xlsx_sheet_parts(parts: dict) -> dict:
    """{sheet name: part path} for worksheets and chartsheets, via workbook.xml and its rels."""
    from lxml import etree
    wb = etree.fromstring(parts["xl/workbook.xml"])
    rels = etree.fromstring(parts["xl/_rels/workbook.xml.rels"])
    targets = {r.get("Id"): r.get("Target") for r in rels}
    out = {}
    for sh in wb.iter(f"{{{_NS_MAIN}}}sheet"):
        target = targets.get(sh.get(f"{{{_NS_REL}}}id"), "")
        out[sh.get("name")] = target.lstrip("/") if target.startswith("/") else "xl/" + target
    return out


def watermark_xlsx(data: bytes, ctx: WatermarkContext, manifest: Optional[dict] = None) -> bytes:
    """Package-level XLSX watermark. Every part of the workbook (cells, formulas,
    styles, charts, pivot tables, images, defined names, VBA-free extensions) is
    copied byte-for-byte; only these changes are made:
      * each worksheet/chartsheet gets a header (identity + action ID, in red) and a
        footer (CONFIDENTIAL + action ID + page numbers), visible in Page Layout/print;
      * a "LexCorp Control" worksheet is appended with the full watermark block, the
        cumulative modification history and the verification QR image.
    Cells are never overlaid, so the workbook stays usable."""
    from lxml import etree

    zin = zipfile.ZipFile(io.BytesIO(data))
    names = zin.namelist()
    parts = {n: zin.read(n) for n in names}

    # Excel limits a header/footer to 255 characters.
    header = ("&C&8&KC00000" + _hf_escape(ctx.compact()))[:250]
    footer = f"&C&8CONFIDENTIAL - {_hf_escape(ctx.action_id)} - Page &P of &N"
    original_hf = {}
    for sheet_name, part in xlsx_sheet_parts(parts).items():
        if part in parts:
            existing_hf = etree.fromstring(parts[part]).find(f"{{{_NS_MAIN}}}headerFooter")
            original_hf[sheet_name] = etree.tostring(existing_hf).decode() if existing_hf is not None else None
            parts[part] = _xlsx_set_header_footer(parts[part], header, footer)

    # ---- workbook: register the control sheet
    wb = etree.fromstring(parts["xl/workbook.xml"])
    sheets = wb.find(f"{{{_NS_MAIN}}}sheets")
    existing = {sh.get("name") for sh in sheets}
    title = "LexCorp Control"
    while title in existing:
        title += "_"
    sheet_id = max([int(sh.get("sheetId")) for sh in sheets] + [0]) + 1
    rels_path = "xl/_rels/workbook.xml.rels"
    rels = etree.fromstring(parts[rels_path])
    rid = "rIdLexCorpControl"
    while any(r.get("Id") == rid for r in rels):
        rid += "x"
    n = 1
    while f"xl/worksheets/lexcorpControl{n}.xml" in parts:
        n += 1
    ws_part = f"xl/worksheets/lexcorpControl{n}.xml"
    etree.SubElement(sheets, f"{{{_NS_MAIN}}}sheet", {"name": title, "sheetId": str(sheet_id), f"{{{_NS_REL}}}id": rid})
    etree.SubElement(rels, f"{{{_NS_PKG_REL}}}Relationship",
                     {"Id": rid, "Type": _REL_WS, "Target": f"worksheets/lexcorpControl{n}.xml"})
    parts["xl/workbook.xml"] = etree.tostring(wb, xml_declaration=True, encoding="UTF-8", standalone=True)
    parts[rels_path] = etree.tostring(rels, xml_declaration=True, encoding="UTF-8", standalone=True)

    # ---- the control worksheet (inline strings: sharedStrings is left untouched)
    ws = etree.Element(f"{{{_NS_MAIN}}}worksheet", nsmap={None: _NS_MAIN, "r": _NS_REL})
    cols = etree.SubElement(ws, f"{{{_NS_MAIN}}}cols")
    etree.SubElement(cols, f"{{{_NS_MAIN}}}col", {"min": "1", "max": "1", "width": "70", "customWidth": "1"})
    data_el = etree.SubElement(ws, f"{{{_NS_MAIN}}}sheetData")
    lines = [CONTROL_TITLE, ""] + ctx.all_lines()
    for i, line in enumerate(lines, start=1):
        row = etree.SubElement(data_el, f"{{{_NS_MAIN}}}row", {"r": str(i)})
        if line:
            c = etree.SubElement(row, f"{{{_NS_MAIN}}}c", {"r": f"A{i}", "t": "inlineStr"})
            is_ = etree.SubElement(c, f"{{{_NS_MAIN}}}is")
            etree.SubElement(is_, f"{{{_NS_MAIN}}}t").text = line
    _xlsx_hf_root = etree.SubElement(ws, f"{{{_NS_MAIN}}}headerFooter")
    etree.SubElement(_xlsx_hf_root, f"{{{_NS_MAIN}}}oddHeader").text = header
    drawing_rid = "rIdLexCorpQr"
    etree.SubElement(ws, f"{{{_NS_MAIN}}}drawing", {f"{{{_NS_REL}}}id": drawing_rid})
    parts[ws_part] = etree.tostring(ws, xml_declaration=True, encoding="UTF-8", standalone=True)

    # ---- QR image via a drawing part
    media = f"xl/media/lexcorpQr{n}.png"
    drawing = f"xl/drawings/lexcorpControl{n}.xml"
    parts[media] = ctx.qr_png()
    parts[f"xl/worksheets/_rels/lexcorpControl{n}.xml.rels"] = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<Relationships xmlns="{_NS_PKG_REL}"><Relationship Id="{drawing_rid}" Type="{_REL_DRAWING}" '
        f'Target="../drawings/lexcorpControl{n}.xml"/></Relationships>').encode()
    parts[f"xl/drawings/_rels/lexcorpControl{n}.xml.rels"] = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<Relationships xmlns="{_NS_PKG_REL}"><Relationship Id="rIdQr" Type="{_REL_IMAGE}" '
        f'Target="../media/lexcorpQr{n}.png"/></Relationships>').encode()
    emu = 1_524_000  # 160 px
    parts[drawing] = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<xdr:wsDr xmlns:xdr="{_NS_XDR}" xmlns:a="{_NS_A}" xmlns:r="{_NS_REL}">'
        '<xdr:oneCellAnchor><xdr:from><xdr:col>2</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>2</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from>'
        f'<xdr:ext cx="{emu}" cy="{emu}"/>'
        '<xdr:pic><xdr:nvPicPr><xdr:cNvPr id="2" name="LexCorp verification QR"/><xdr:cNvPicPr/></xdr:nvPicPr>'
        '<xdr:blipFill><a:blip r:embed="rIdQr"/><a:stretch><a:fillRect/></a:stretch></xdr:blipFill>'
        f'<xdr:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{emu}" cy="{emu}"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></xdr:spPr>'
        '</xdr:pic><xdr:clientData/></xdr:oneCellAnchor></xdr:wsDr>').encode()

    ct = etree.fromstring(parts["[Content_Types].xml"])
    if not any(d.get("Extension", "").lower() == "png" for d in ct.findall(f"{{{_NS_CT}}}Default")):
        etree.SubElement(ct, f"{{{_NS_CT}}}Default", {"Extension": "png", "ContentType": "image/png"})
    etree.SubElement(ct, f"{{{_NS_CT}}}Override", {"PartName": "/" + ws_part,
                     "ContentType": "application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"})
    etree.SubElement(ct, f"{{{_NS_CT}}}Override", {"PartName": "/" + drawing,
                     "ContentType": "application/vnd.openxmlformats-officedocument.drawing+xml"})
    parts["[Content_Types].xml"] = etree.tostring(ct, xml_declaration=True, encoding="UTF-8", standalone=True)
    if manifest is not None:
        manifest.update({"format": "xlsx", "header": header, "footer": footer, "original_headers": original_hf,
                         "control_sheet": title, "control_lines": [x for x in lines if x], "qr": sha256_hex(parts[media])})

    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        order = ["[Content_Types].xml"] + [x for x in names if x != "[Content_Types].xml"] + \
                [x for x in parts if x not in names and x != "[Content_Types].xml"]
        for x in order:
            zout.writestr(x, parts[x])
    return out.getvalue()


# ------------------------------------------------------------------ entry point

_MIME = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


def render_protected_copy_with_manifest(data: bytes, filename: str, ctx: WatermarkContext) -> tuple[bytes, str, str, dict]:
    """(watermarked_bytes, mime_type, output_filename, manifest). The manifest fingerprints
    every watermark element so the copy can later be verified and cleanly un-watermarked."""
    kind = detect_type(data, filename)
    stem = filename.rsplit(".", 1)[0] if "." in filename else filename
    tag = f"{stem}.{ctx.action_id}"
    m: dict = {"action_id": ctx.action_id, "qr_url": ctx.verify_url}
    if kind == "pdf":
        return watermark_pdf(data, ctx, m), _MIME["pdf"], f"{tag}.pdf", m
    if kind == "text":
        return watermark_pdf(text_to_pdf(data, filename), ctx, m), _MIME["pdf"], f"{tag}.pdf", m
    if kind == "image":
        out, mime = watermark_image(data, ctx, m)
        return out, mime, f"{tag}.{'jpg' if mime == 'image/jpeg' else 'png'}", m
    if kind == "docx":
        return watermark_docx(data, ctx, m), _MIME["docx"], f"{tag}.docx", m
    if kind == "pptx":
        return watermark_pptx(data, ctx, m), _MIME["pptx"], f"{tag}.pptx", m
    if kind == "xlsx":
        return watermark_xlsx(data, ctx, m), _MIME["xlsx"], f"{tag}.xlsx", m
    raise UnsupportedDocumentType(kind)


def render_protected_copy(data: bytes, filename: str, ctx: WatermarkContext) -> tuple[bytes, str, str]:
    """Return (watermarked_bytes, mime_type, output_filename). Never returns the original."""
    content, mime, name, _ = render_protected_copy_with_manifest(data, filename, ctx)
    return content, mime, name


# ------------------------------------------------------------------ view-only representation

def office_text(data: bytes, kind: str) -> str:
    """Readable text content of an Office document (for the browser VIEW rendering)."""
    out: List[str] = []
    if kind == "docx":
        import docx
        d = docx.Document(io.BytesIO(data))
        out.extend(p.text for p in d.paragraphs)
        for t_i, table in enumerate(d.tables, start=1):
            out.append(f"")
            out.append(f"[Table {t_i}]")
            for row in table.rows:
                out.append(" | ".join(c.text.strip() for c in row.cells))
    elif kind == "pptx":
        from pptx import Presentation
        prs = Presentation(io.BytesIO(data))
        for i, slide in enumerate(prs.slides, start=1):
            out.append(f"--- Slide {i} ---")
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                    out.extend(p.text for p in shape.text_frame.paragraphs if p.text.strip())
            out.append("")
    elif kind == "xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=False, read_only=True)
        for ws in wb.worksheets:
            out.append(f"--- Sheet: {ws.title} ---")
            for row in ws.iter_rows(values_only=True):
                if any(v is not None for v in row):
                    out.append(" | ".join("" if v is None else str(v) for v in row))
            out.append("")
    return "\n".join(out)


def render_view_copy(data: bytes, filename: str, ctx: WatermarkContext) -> tuple[bytes, str, str]:
    """VIEW representation for the browser: always a watermarked PDF or image, never the
    document's native Office format (a native copy would be a download in disguise;
    full-fidelity copies go through the controlled download workflow). Office files are
    rendered as a watermarked PDF of their text content; layout is not reproduced."""
    kind = detect_type(data, filename)
    if kind in ("docx", "pptx", "xlsx"):
        stem = filename.rsplit(".", 1)[0] if "." in filename else filename
        text = f"View rendering of {filename} (text content only; layout not reproduced)\n\n" + office_text(data, kind)
        return watermark_pdf(text_to_pdf(text.encode("utf-8"), filename), ctx), _MIME["pdf"], f"{stem}.{ctx.action_id}.pdf"
    return render_protected_copy(data, filename, ctx)
