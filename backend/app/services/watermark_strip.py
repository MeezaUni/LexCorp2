"""Verify, then remove, LexCorp's own watermark layers from a returned copy.

watermark.render_protected_copy_with_manifest() records a MANIFEST for every copy it
issues: fingerprints of each watermark element. strip_watermark() takes the returned file
and that manifest and, for every element:
  * checks it is still there and unchanged (a changed or missing element is a FINDING,
    with the page / slide / sheet where it was found);
  * removes it, together with the copy's control page/section/slide/sheet.
The result is the person's edit without any LexCorp watermark, which becomes the new
version; later copies of it get one fresh watermark instead of a stack.

Findings: severity "tamper" (a watermark element was changed, removed or could not be
processed), or "review" (something unusual that may be a legitimate edit, e.g. a page
that carried a watermark was deleted). Nothing here is shown to the uploader.
"""

import io
import re
import zipfile
from typing import Optional

from app.services import hidden_mark
from app.services.watermark import (
    CONTROL_KEY, LAYER_KEY, _AFTER_HF, _NS_MAIN, _NS_PKG_REL, _NS_REL, _NS_CT, _W_NS,
    detect_type, image_band, image_overlay_layer, sha256_hex, xlsx_sheet_parts,
)


def finding(severity: str, where: str, detail: str, page: Optional[int] = None) -> dict:
    out = {"severity": severity, "where": where, "detail": detail}
    if page is not None:
        out["page"] = page
    return out


def strip_watermark(data: bytes, filename: str, manifest: dict) -> tuple[bytes, list[dict]]:
    """(clean bytes, findings). On any failure the original bytes are returned with a finding."""
    fmt = manifest.get("format")
    try:
        kind = detect_type(data, filename or "")
    except Exception:
        kind = None
    expected_kind = "pdf" if fmt == "pdf" else fmt
    if kind != expected_kind:
        return data, [finding("tamper", "whole file", f"The copy was issued as {fmt or 'unknown'} but a {kind or 'unrecognised'} "
                                                      "file came back, so its watermark could not be verified or removed.")]
    try:
        clean, findings = {"pdf": _strip_pdf, "docx": _strip_docx, "pptx": _strip_pptx, "xlsx": _strip_xlsx,
                           "image": _strip_image}[fmt](data, manifest)
        if fmt in ("docx", "pptx", "xlsx"):
            clean = hidden_mark.remove_ooxml(clean, manifest)
        return clean, findings
    except Exception as e:  # never let a malformed file pass silently
        return data, [finding("tamper", "whole file", f"The watermark could not be processed ({type(e).__name__}); "
                                                      "the file may have been altered.")]


# ------------------------------------------------------------------ PDF

_INVOKE = re.compile(rb"q\s+1 0 0 1 [-\d.]+ [-\d.]+ cm\s*/LCWM\d+\s+Do\s+Q")


def _form_images(form) -> list[str]:
    res = form.get("/Resources")
    xo = res.get_object().get("/XObject") if res is not None else None
    if xo is None:
        return []
    xo = xo.get_object()
    return sorted(sha256_hex(xo[k].get_object().get_data()) for k in xo)


def _strip_pdf(data: bytes, m: dict):
    from pypdf import PdfReader, PdfWriter
    from pypdf.generic import ArrayObject, DecodedStreamObject, NameObject

    aid = m["action_id"]
    expected = m.get("overlays", {})
    findings, seen, control_idx, control_found = [], set(), [], []
    reader = PdfReader(io.BytesIO(data))
    writer = PdfWriter(clone_from=reader)
    for idx, page in enumerate(writer.pages, start=1):
        tag = page.get(CONTROL_KEY)
        if tag is not None:
            control_idx.append(idx - 1)
            if str(tag) == aid:
                c = page.get_contents()
                control_found.append((idx, sha256_hex(c.get_data() if c is not None else b"")))
            continue
        res = page.get("/Resources")
        xo = res.get_object().get("/XObject") if res is not None else None
        xo = xo.get_object() if xo is not None else None
        for key in list(xo.keys()) if xo is not None else []:
            obj = xo[key].get_object()
            name = str(key)[1:]
            if obj.get(LAYER_KEY) is None and not name.startswith("LCWM"):
                continue
            if str(obj.get(LAYER_KEY)) == aid and name in expected:
                info = expected[name]
                if sha256_hex(obj.get_data()) != info["form"] or _form_images(obj) != info["images"]:
                    findings.append(finding("tamper", f"page {idx}", "The watermark layer on this page was changed.", idx))
                seen.add(name)
            del xo[key]
        contents = page.get("/Contents")
        if contents is not None:
            c = contents.get_object()
            if isinstance(c, ArrayObject):
                page[NameObject("/Contents")] = ArrayObject([x for x in c if x.get_object().get(LAYER_KEY) is None])
            raw = page.get_contents().get_data() if page.get_contents() is not None else b""
            cleaned = _INVOKE.sub(b"", raw)
            if cleaned != raw:
                st = DecodedStreamObject()
                st.set_data(cleaned)
                page[NameObject("/Contents")] = writer._add_object(st)

    body_pages = len(writer.pages) - len(control_idx)
    for name, info in sorted(expected.items(), key=lambda kv: kv[1]["page"]):
        if name in seen:
            continue
        if body_pages < m.get("page_count", 0):
            findings.append(finding("review", f"original page {info['page']}",
                                    "This page and its watermark are no longer in the file (page deleted?)."))
        else:
            findings.append(finding("tamper", f"original page {info['page']}",
                                    "The watermark was removed from this page.", info["page"]))
    if not control_found:
        findings.append(finding("tamper", "Document Control page", "The Document Control page was removed."))
    elif [h for _, h in control_found] != m.get("control_pages", []):
        idx = control_found[0][0]
        findings.append(finding("tamper", f"page {idx} (Document Control)", "The Document Control page was changed.", idx))
    for i in reversed(control_idx):
        del writer.pages[i]
    writer.compress_identical_objects(remove_identicals=False, remove_orphans=True)  # drop the unreferenced layers
    for key in ("/Subject", "/Keywords"):
        if writer._info is not None and key in writer._info.get_object():
            del writer._info.get_object()[key]
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue(), findings


# ------------------------------------------------------------------ DOCX

def _w(tag: str) -> str:
    return f"{{{_W_NS}}}{tag}"


def _text(el) -> str:
    return "".join(t.text or "" for t in el.iter(_w("t")))


def _embedded_images(part, el) -> list:
    ids = [v for node in el.iter() for k, v in node.attrib.items() if k.endswith("}embed") or k.endswith("}id") and v.startswith("rId")]
    out = []
    for rid in ids:
        rel = part.rels.get(rid)
        if rel is not None and not rel.is_external and rel.reltype.endswith("/image"):
            out.append((rid, sha256_hex(rel.target_part.blob)))
    return out


def _strip_docx(data: bytes, m: dict):
    import docx

    aid = m["action_id"]
    findings, headers_found, control_found = [], 0, False
    doc = docx.Document(io.BytesIO(data))
    parts = [p for p in doc.part.package.iter_parts()
             if hasattr(p, "_element") and re.match(r"/word/(document|header\d*|footer\d*)\.xml$", str(p.partname))]
    for part in parts:
        where = "document body" if "document" in str(part.partname) else f"header/footer {str(part.partname).rsplit('/', 1)[-1]}"
        for sdt in list(part._element.iter(_w("sdt"))):
            tag_el = sdt.find(f"{_w('sdtPr')}/{_w('tag')}")
            tag = tag_el.get(_w("val")) if tag_el is not None else ""
            if not tag.startswith(("LexCorpWM:", "LexCorpControl:")):
                continue
            images = _embedded_images(part, sdt)
            if tag == f"LexCorpWM:{aid}":
                headers_found += 1
                shape = next((n for n in sdt.iter() if isinstance(n.tag, str) and n.tag.endswith("}textpath")), None)
                ok = (_text(sdt).strip() == m["header_text"] and shape is not None and shape.get("string") == m["shape_string"]
                      and [h for _, h in images] == [m["qr"]])
                if not ok:
                    findings.append(finding("tamper", where, "The watermark in this header was changed."))
            elif tag == f"LexCorpControl:{aid}":
                control_found = True
                lines = [_text(p) for p in sdt.iter(_w("p")) if _text(p)]
                if lines != m["control_lines"] or [h for _, h in images] != [m["qr"]]:
                    findings.append(finding("tamper", "Document Control section", "The Document Control section was changed."))
            sdt.getparent().remove(sdt)
            for rid, _ in images:
                part.drop_rel(rid)
        # Watermark text left behind outside its content control (the control was removed)
        for p in list(part._element.iter(_w("p"))):
            t = _text(p)
            if aid in t or t.strip() == m["control_lines"][0]:
                findings.append(finding("tamper", where, "Watermark text was found outside its protected layer."))
                p.getparent().remove(p)
    if headers_found < m.get("headers", 0):
        findings.append(finding("tamper", "headers", f"The watermark was removed from {m['headers'] - headers_found} header(s)."))
    if not control_found:
        findings.append(finding("tamper", "Document Control section", "The Document Control section was removed."))
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue(), findings


# ------------------------------------------------------------------ PPTX

def _strip_pptx(data: bytes, m: dict):
    from pptx import Presentation

    aid = m["action_id"]
    findings, marked, control_slides, control_ok = [], set(), [], None
    prs = Presentation(io.BytesIO(data))
    slides = list(prs.slides)
    for sidx, slide in enumerate(slides, start=1):
        is_control = False
        for shape in list(slide.shapes):
            name = shape.name or ""
            if name.startswith("LexCorpControl"):
                is_control = True
                if name == f"LexCorpControl:{aid}":
                    lines = [p.text for p in shape.text_frame.paragraphs if p.text] if shape.has_text_frame else []
                    control_ok = (control_ok is not False) and lines == m["control_lines"]
                elif name == f"LexCorpControlQR:{aid}":
                    control_ok = (control_ok is not False) and sha256_hex(shape.image.blob) == m["qr"]
                continue
            if not name.startswith(("LexCorpWM:", "LexCorpWMQR:")):
                continue
            bits = name.split(":")
            if len(bits) == 3 and bits[1] == aid:
                if bits[0] == "LexCorpWMQR":
                    if sha256_hex(shape.image.blob) != m["qr"]:
                        findings.append(finding("tamper", f"slide {sidx}", "The watermark QR code on this slide was changed.", sidx))
                else:
                    marked.add(bits[2])
                    if not shape.has_text_frame or shape.text_frame.text != m["box_text"]:
                        findings.append(finding("tamper", f"slide {sidx}", "The watermark on this slide was changed.", sidx))
            shape._element.getparent().remove(shape._element)
        if is_control:
            control_slides.append(slide)
    missing = [n for n in range(1, m.get("slide_count", 0) + 1) if str(n) not in marked]
    if missing:
        fewer = len(slides) - len(control_slides) < m.get("slide_count", 0)
        for n in missing:
            findings.append(finding("review" if fewer else "tamper", f"original slide {n}",
                                    "This slide and its watermark are no longer in the file (slide deleted?)." if fewer
                                    else "The watermark was removed from this slide.", None if fewer else n))
    if control_ok is None:
        findings.append(finding("tamper", "Document Control slide", "The Document Control slide was removed."))
    elif control_ok is False:
        findings.append(finding("tamper", "Document Control slide", "The Document Control slide was changed."))
    sld_ids = prs.slides._sldIdLst
    for slide in control_slides:
        for sld in list(sld_ids):
            if prs.part.related_part(sld.rId) is slide.part:
                prs.part.drop_rel(sld.rId)
                sld_ids.remove(sld)
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue(), findings


# ------------------------------------------------------------------ XLSX

def _rels_path(part: str) -> str:
    d, f = part.rsplit("/", 1)
    return f"{d}/_rels/{f}.rels"


def _resolve(base_part: str, target: str) -> str:
    if target.startswith("/"):
        return target.lstrip("/")
    parts = base_part.rsplit("/", 1)[0].split("/")
    for seg in target.split("/"):
        if seg == "..":
            parts.pop()
        elif seg != ".":
            parts.append(seg)
    return "/".join(parts)


def _strip_xlsx(data: bytes, m: dict):
    from lxml import etree

    findings = []
    zin = zipfile.ZipFile(io.BytesIO(data))
    names = zin.namelist()
    parts = {n: zin.read(n) for n in names}
    sheets = xlsx_sheet_parts(parts)
    shared = []
    if "xl/sharedStrings.xml" in parts:
        sst = etree.fromstring(parts["xl/sharedStrings.xml"])
        shared = ["".join(t.text or "" for t in si.iter(f"{{{_NS_MAIN}}}t")) for si in sst.iter(f"{{{_NS_MAIN}}}si")]

    removed = set()
    ctl_name = m["control_sheet"]
    ctl_part = sheets.pop(ctl_name, None)
    if ctl_part is None or ctl_part not in parts:
        findings.append(finding("tamper", "LexCorp Control sheet", "The control sheet was removed."))
    else:
        ws = etree.fromstring(parts[ctl_part])
        lines = []
        for c in ws.iter(f"{{{_NS_MAIN}}}c"):
            if c.get("t") == "s":
                v = c.find(f"{{{_NS_MAIN}}}v")
                lines.append(shared[int(v.text)] if v is not None and v.text and int(v.text) < len(shared) else "")
            else:
                lines.append("".join(t.text or "" for t in c.iter(f"{{{_NS_MAIN}}}t")))
        lines = [x for x in lines if x]
        qr_hashes = []
        to_remove = [ctl_part]
        rels = _rels_path(ctl_part)
        if rels in parts:
            to_remove.append(rels)
            for r in etree.fromstring(parts[rels]):
                if r.get("Type", "").endswith("/drawing"):
                    drawing = _resolve(ctl_part, r.get("Target"))
                    to_remove.append(drawing)
                    drels = _rels_path(drawing)
                    if drels in parts:
                        to_remove.append(drels)
                        for dr in etree.fromstring(parts[drels]):
                            if dr.get("Type", "").endswith("/image"):
                                media = _resolve(drawing, dr.get("Target"))
                                if media in parts:
                                    qr_hashes.append(sha256_hex(parts[media]))
                                    to_remove.append(media)
        if lines != m["control_lines"] or qr_hashes != [m["qr"]]:
            findings.append(finding("tamper", "LexCorp Control sheet", "The control sheet was changed."))
        removed.update(p for p in to_remove if p in parts)

        wb = etree.fromstring(parts["xl/workbook.xml"])
        sheets_el = wb.find(f"{{{_NS_MAIN}}}sheets")
        position, rid = None, None
        for i, sh in enumerate(list(sheets_el)):
            if sh.get("name") == ctl_name:
                position, rid = i, sh.get(f"{{{_NS_REL}}}id")
                sheets_el.remove(sh)
        dn = wb.find(f"{{{_NS_MAIN}}}definedNames")
        if dn is not None and position is not None:
            for d in list(dn):
                local = d.get("localSheetId")
                if local is not None and int(local) == position:
                    dn.remove(d)
                elif local is not None and int(local) > position:
                    d.set("localSheetId", str(int(local) - 1))
        parts["xl/workbook.xml"] = etree.tostring(wb, xml_declaration=True, encoding="UTF-8", standalone=True)
        wrels = etree.fromstring(parts["xl/_rels/workbook.xml.rels"])
        for r in list(wrels):
            if r.get("Id") == rid:
                wrels.remove(r)
        parts["xl/_rels/workbook.xml.rels"] = etree.tostring(wrels, xml_declaration=True, encoding="UTF-8", standalone=True)
        ct = etree.fromstring(parts["[Content_Types].xml"])
        for o in list(ct.findall(f"{{{_NS_CT}}}Override")):
            if o.get("PartName", "").lstrip("/") in removed:
                ct.remove(o)
        parts["[Content_Types].xml"] = etree.tostring(ct, xml_declaration=True, encoding="UTF-8", standalone=True)

    # Headers/footers: verify ours, then put back what each sheet had before
    original = m.get("original_headers", {})
    for name, part in sheets.items():
        if part not in parts or name not in original:
            continue
        root = etree.fromstring(parts[part])
        hf = root.find(f"{{{_NS_MAIN}}}headerFooter")
        head = hf.find(f"{{{_NS_MAIN}}}oddHeader") if hf is not None else None
        foot = hf.find(f"{{{_NS_MAIN}}}oddFooter") if hf is not None else None
        if head is None or head.text != m["header"] or foot is None or foot.text != m["footer"]:
            findings.append(finding("tamper", f"sheet '{name}'", "The watermark header/footer on this sheet was changed or removed."))
        if hf is not None:
            root.remove(hf)
        if original[name]:
            restored = etree.fromstring(original[name])
            anchor = next((c for c in root if isinstance(c.tag, str) and etree.QName(c).localname in _AFTER_HF), None)
            if anchor is not None:
                anchor.addprevious(restored)
            else:
                root.append(restored)
        parts[part] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    for name in original:
        if name != ctl_name and name not in sheets:
            findings.append(finding("review", f"sheet '{name}'", "This sheet is no longer in the workbook (deleted or renamed?)."))

    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in ["[Content_Types].xml"] + [x for x in names if x != "[Content_Types].xml"]:
            if n not in removed:
                zout.writestr(n, parts[n])
    return out.getvalue(), findings


# ------------------------------------------------------------------ images

def _strip_image(data: bytes, m: dict):
    import numpy as np
    from PIL import Image, ImageChops
    from app.services.qr import generate_qr_for_url

    img = Image.open(io.BytesIO(data)).convert("RGBA")
    w, h, bh = m["width"], m["height"], m["band_height"]
    if img.size != (w, h + bh):
        return data, [finding("tamper", "whole image", f"The image size changed ({img.size[0]}x{img.size[1]}, issued as "
                                                       f"{w}x{h + bh}), so its watermark could not be verified or removed.")]
    findings = []
    band = img.crop((0, h, w, h + bh)).convert("RGB")
    expected = image_band(w, m["band_lines"], generate_qr_for_url(m["qr_url"]), m.get("mark")).convert("RGB")
    diff = ImageChops.difference(band, expected).convert("L").point(lambda v: 255 if v > (60 if m.get("jpeg") else 24) else 0)
    box = diff.getbbox()
    if box:
        findings.append(finding("tamper", f"footer band (x {box[0]}-{box[2]}, y {box[1]}-{box[3]})",
                                "The watermark band below the image was changed."))
    # Reverse the semi-transparent marking: original = (marked - a*C) / (1 - a)
    layer = np.asarray(image_overlay_layer(w, h, m["overlay_text"], m["overlay_size"]), dtype=np.float64)
    a = layer[..., 3:4] / 255.0
    marked = np.asarray(img.crop((0, 0, w, h)), dtype=np.float64)
    rgb = (marked[..., :3] - a * layer[..., :3]) / np.clip(1.0 - a, 1e-6, None)
    clean = np.dstack([np.clip(rgb, 0, 255), marked[..., 3:4]]).round().astype(np.uint8)
    out = io.BytesIO()
    result = Image.fromarray(clean, "RGBA")
    if m.get("jpeg"):
        result.convert("RGB").save(out, format="JPEG", quality=95)
    else:
        result.save(out, format="PNG")
    return out.getvalue(), findings
